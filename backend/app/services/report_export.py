"""
Report export (§28) — แปลง report dict → ไฟล์ Excel / PDF (bytes)

import openpyxl/reportlab แบบ lazy ในฟังก์ชัน เพื่อให้ module import ได้
แม้ lib ยังไม่ติดตั้ง (endpoint จะคืน 503 พร้อมข้อความชัดเจน)
"""
from __future__ import annotations

import io


class ExportUnavailable(Exception):
    """lib export ไม่พร้อม (ยังไม่ pip install) — endpoint แปลงเป็น 503"""


def rows_to_xlsx(sheet_title: str, headers: list[str], rows: list[list]) -> bytes:
    """สร้าง Excel 1 sheet จาก headers + rows → bytes"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as e:  # pragma: no cover
        raise ExportUnavailable("openpyxl ยังไม่ติดตั้ง — รัน pip install openpyxl") from e

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Report"  # Excel จำกัด 31 ตัว
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    # auto-width คร่าวๆ
    for col_idx, header in enumerate(headers, start=1):
        width = max(len(str(header)), 12) + 2
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_to_pdf(title: str, sections: list[tuple[str, list[str], list[list]]]) -> bytes:
    """
    สร้าง PDF จากหลาย section — แต่ละ section = (heading, headers, rows)
    ใช้ reportlab (pure-python ไม่ต้อง system lib)
    """
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    except ImportError as e:  # pragma: no cover
        raise ExportUnavailable("reportlab ยังไม่ติดตั้ง — รัน pip install reportlab") from e

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title=title)
    styles = getSampleStyleSheet()
    flow = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

    for heading, headers, rows in sections:
        flow.append(Paragraph(heading, styles["Heading2"]))
        flow.append(Spacer(1, 6))
        data = [headers] + ([[str(c) for c in r] for r in rows] or [["(ไม่มีข้อมูล)"] + [""] * (len(headers) - 1)])
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        flow.append(table)
        flow.append(Spacer(1, 16))

    doc.build(flow)
    return buf.getvalue()
